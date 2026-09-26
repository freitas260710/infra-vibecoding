"""Testes de importar e exportar planilha na tela de banco (US I.1, D53) e da coluna Acesso dos usuários."""
import csv
import io
import re
from datetime import timedelta

import openpyxl
import pytest
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.utils import timezone

from infra_vibecoding import planilhas
from infra_vibecoding.admin import situacao_do_acesso
from infra_vibecoding.login import enviar_link_de_senha
from tests.app_teste.models import Pedido

pytestmark = pytest.mark.django_db
Usuario = get_user_model()
SENHA = "uma-senha-bem-longa-para-teste"
LISTA = "/gestao-interna/app_teste/pedido/"
USUARIOS = "/gestao-interna/app_teste/usuarioteste/"


@pytest.fixture
def chefe():
    return Usuario.objects.create_superuser("chefe@exemplo.com", password=SENHA)


@pytest.fixture
def admin(client, chefe):
    client.force_login(chefe)
    return client


@pytest.fixture
def ana():
    return Usuario.objects.create_user("ana@exemplo.com", password=SENHA, nome="Ana")


def pedido(dono, titulo, valor="10.00"):
    p = Pedido(dono=dono, titulo=titulo, valor=valor)
    p.salvar_como_sistema("teste")
    return p


def planilha_csv(*linhas, separador=";"):
    saida = io.StringIO()
    csv.writer(saida, delimiter=separador).writerows(linhas)
    return ("﻿" + saida.getvalue()).encode("utf-8")


def planilha_xlsx(*linhas):
    livro = openpyxl.Workbook()
    for linha in linhas:
        livro.active.append(list(linha))
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def enviar(client, url, dados, nome="dados.csv"):
    return client.post(url + "importar/", {"arquivo": SimpleUploadedFile(nome, dados)})


def token(resposta):
    return re.search(r'name="token" value="([^"]+)"', resposta.content.decode()).group(1)


def confirmar(client, url, resposta):
    return client.post(url + "importar/", {"token": token(resposta), "confirmar": "1"})


def ler_csv(resposta):
    texto = resposta.content.decode("utf-8")
    assert texto.startswith("﻿")
    return list(csv.reader(io.StringIO(texto[1:]), delimiter=";"))


# 1. Exportar
def test_lista_mostra_os_botoes(admin):
    tela = admin.get(LISTA).content.decode()
    assert "Importar planilha" in tela and "Exportar CSV" in tela and "Exportar Excel" in tela


def test_exporta_csv_no_padrao_do_excel_brasileiro(admin, ana, caplog):
    pedido(ana, "Cadeira", "1234.50")
    linhas = ler_csv(admin.get(LISTA + "exportar/csv/"))
    assert linhas[0] == ["id", "dono", "titulo", "valor"]
    assert linhas[1][1:] == [str(ana.pk), "Cadeira", "1234,50"]
    assert "exportou 1 linha(s) de app_teste.Pedido (csv)" in caplog.text


def test_exporta_so_o_que_a_busca_e_os_filtros_mostram(admin, ana):
    pedido(ana, "Cadeira")
    pedido(ana, "Mesa")
    linhas = ler_csv(admin.get(LISTA + "exportar/csv/?q=Mesa"))
    assert [l[2] for l in linhas[1:]] == ["Mesa"]


def test_exportacao_protegida_contra_formula(admin, ana):
    pedido(ana, '=HYPERLINK("http://mal.exemplo","clique")')
    linhas = ler_csv(admin.get(LISTA + "exportar/csv/"))
    assert linhas[1][2].startswith("'=")
    livro = openpyxl.load_workbook(io.BytesIO(admin.get(LISTA + "exportar/xlsx/").content))
    celula = livro.active.cell(row=2, column=3)
    assert celula.data_type == "s" and celula.value.startswith("'=")


def test_exporta_excel(admin, ana):
    pedido(ana, "Cadeira", "5.00")
    r = admin.get(LISTA + "exportar/xlsx/")
    assert r["Content-Disposition"].endswith('.xlsx"')
    linhas = list(openpyxl.load_workbook(io.BytesIO(r.content)).active.values)
    assert linhas[0] == ("id", "dono", "titulo", "valor") and linhas[1][2] == "Cadeira"


def test_usuarios_exportados_sem_senha_nem_segredos(admin, ana):
    cabecalho = ler_csv(admin.get(USUARIOS + "exportar/csv/"))[0]
    for proibido in ("password", "chave_de_sessao", "segredo_do_app", "codigos_de_recuperacao",
                     "ultimo_codigo_do_app"):
        assert proibido not in cabecalho
    assert "email" in cabecalho


def test_exportar_formato_desconhecido(admin):
    assert admin.get(LISTA + "exportar/pdf/").status_code == 404


# 2. Importar: prévia, tudo ou nada, criar e atualizar
def test_previa_nao_grava_e_confirmar_grava(admin, ana, caplog):
    dados = planilha_csv(["dono", "titulo", "valor"], [ana.pk, "Cadeira", "1.234,50"], [ana.pk, "Mesa", "10"])
    r = enviar(admin, LISTA, dados)
    assert r.status_code == 200 and "2</strong> nova(s)" in r.content.decode()
    assert Pedido.objects.como_sistema("teste").count() == 0
    r = confirmar(admin, LISTA, r)
    assert r.status_code == 302
    valores = sorted(Pedido.objects.como_sistema("teste").values_list("titulo", "valor"))
    assert [(t, str(v)) for t, v in valores] == [("Cadeira", "1234.50"), ("Mesa", "10.00")]
    assert "importou a planilha 'dados.csv' em app_teste.Pedido: 2 novo(s), 0 atualizado(s)" in caplog.text


def test_tudo_ou_nada(admin, ana):
    dados = planilha_csv(["dono", "titulo", "valor"], [ana.pk, "Cadeira", "10"], [ana.pk, "", "abc"])
    r = enviar(admin, LISTA, dados)
    tela = r.content.decode()
    assert "Nada foi gravado" in tela and "Confirmar e gravar" not in tela
    assert "<td>3</td><td>titulo</td>" in tela and "<td>3</td><td>valor</td>" in tela
    assert Pedido.objects.como_sistema("teste").count() == 0


def test_nao_da_para_confirmar_uma_previa_com_erro(admin, ana):
    r = enviar(admin, LISTA, planilha_csv(["dono", "titulo"], [ana.pk, ""]))
    guardado = [k for k in re.findall(r'name="token" value="([^"]+)"', r.content.decode())]
    assert guardado == []  # sem botão de confirmar
    assert Pedido.objects.como_sistema("teste").count() == 0


def test_se_o_banco_recusar_uma_linha_nada_fica_gravado(admin, ana, monkeypatch):
    r = enviar(admin, LISTA, planilha_csv(["dono", "titulo"], [ana.pk, "Um"], [ana.pk, "Dois"], [ana.pk, "Tres"]))
    original = Pedido.salvar_como_sistema

    def falha_na_terceira(self, motivo, *a, **k):
        if self.titulo == "Tres":
            raise RuntimeError("banco recusou")
        return original(self, motivo, *a, **k)

    monkeypatch.setattr(Pedido, "salvar_como_sistema", falha_na_terceira)
    r = confirmar(admin, LISTA, r)
    assert "nada foi gravado" in r.content.decode()
    assert Pedido.objects.como_sistema("teste").count() == 0


def test_atualiza_pelo_id_so_as_colunas_da_planilha(admin, ana):
    p = pedido(ana, "Cadeira", "10.00")
    r = enviar(admin, LISTA, planilha_csv(["id", "titulo"], [p.pk, "Cadeira azul"]))
    assert "1</strong> atualizada(s)" in r.content.decode()
    confirmar(admin, LISTA, r)
    p = Pedido.objects.como_sistema("teste").get(pk=p.pk)
    assert p.titulo == "Cadeira azul" and str(p.valor) == "10.00"


def test_id_que_nao_existe_e_erro(admin, ana):
    r = enviar(admin, LISTA, planilha_csv(["id", "titulo"], [999, "X"]))
    assert "Não existe registro com id 999" in r.content.decode()


def test_criar_sem_coluna_obrigatoria_e_erro(admin):
    r = enviar(admin, LISTA, planilha_csv(["titulo"], ["Sem dono"]))
    assert "precisa da coluna: dono" in r.content.decode()


def test_relacao_por_campo_unico(admin, ana):
    r = enviar(admin, LISTA, planilha_csv(["dono__email", "titulo"], ["ana@exemplo.com", "Cadeira"],
                                          ["ninguem@exemplo.com", "Mesa"]))
    tela = r.content.decode()
    assert "Não encontrado: ninguem@exemplo.com." in tela
    r = enviar(admin, LISTA, planilha_csv(["dono__email", "titulo"], ["ANA@exemplo.com".lower(), "Cadeira"]))
    confirmar(admin, LISTA, r)
    assert Pedido.objects.como_sistema("teste").get().dono == ana


def test_relacao_por_campo_que_nao_e_unico_e_recusada(admin, ana):
    r = enviar(admin, LISTA, planilha_csv(["dono__nome", "titulo"], ["Ana", "Cadeira"]))
    assert "não é um campo único" in r.content.decode()


def test_colunas_desconhecidas_sao_ignoradas_com_aviso(admin, ana):
    r = enviar(admin, LISTA, planilha_csv(["dono", "titulo", "Creation Date"], [ana.pk, "Cadeira", "hoje"]))
    assert '"Creation Date" (não existe nesta tabela)' in r.content.decode()


def test_importa_excel(admin, ana):
    r = enviar(admin, LISTA, planilha_xlsx(["dono", "titulo", "valor"], [ana.pk, "Cadeira", 12.5]), "p.xlsx")
    confirmar(admin, LISTA, r)
    assert str(Pedido.objects.como_sistema("teste").get().valor) == "12.50"


def test_exportar_e_importar_de_volta_so_atualiza(admin, ana):
    pedido(ana, "Cadeira")
    pedido(ana, "=perigo")
    exportado = admin.get(LISTA + "exportar/csv/").content
    r = enviar(admin, LISTA, exportado)
    assert "0</strong> nova(s)" in r.content.decode() and "2</strong> atualizada(s)" in r.content.decode()
    confirmar(admin, LISTA, r)
    assert sorted(Pedido.objects.como_sistema("teste").values_list("titulo", flat=True)) == ["=perigo", "Cadeira"]


# 3. Usuários importados
def test_usuarios_importados_nascem_sem_senha_e_sem_email(admin, mailoutbox):
    dados = planilha_csv(["email", "nome", "is_superuser", "is_staff", "password"],
                         ["bia@exemplo.com", "Bia", "sim", "sim", "123"], ["Caio@Exemplo.com", "Caio", "", "", ""])
    r = enviar(admin, USUARIOS, dados)
    tela = r.content.decode()
    for coluna in ("is_superuser", "is_staff", "password"):
        assert f'"{coluna}" (não é importado por segurança)' in tela
    confirmar(admin, USUARIOS, r)
    bia = Usuario.objects.get_by_natural_key("bia@exemplo.com")
    assert not bia.has_usable_password() and not bia.is_superuser and not bia.is_staff
    assert Usuario.objects.get_by_natural_key("caio@exemplo.com").nome == "Caio"
    assert mailoutbox == []
    assert situacao_do_acesso(bia) == "Link não enviado"


def test_email_repetido_na_planilha_e_erro(admin, ana):
    r = enviar(admin, USUARIOS, planilha_csv(["email"], ["bia@exemplo.com"], ["BIA@exemplo.com"],
                                             ["ana@exemplo.com"]))
    tela = r.content.decode()
    assert "já está na linha 2" in tela
    assert "<td>4</td><td>email</td>" in tela  # ana já existe no banco


# 4. Arquivo recusado inteiro
def test_arquivo_que_nao_e_planilha(admin):
    r = enviar(admin, LISTA, b"MZ\x90\x00\x03\x00\x00\x00binario", "virus.csv")
    assert "não é uma planilha" in r.content.decode()
    r = enviar(admin, LISTA, b"PK\x03\x04lixo", "falso.xlsx")
    assert "não é uma planilha" in r.content.decode()


def test_limite_de_linhas_e_tamanho(admin, ana, monkeypatch):
    monkeypatch.setattr(planilhas, "MAX_LINHAS", 2)
    r = enviar(admin, LISTA, planilha_csv(["dono", "titulo"], *[[ana.pk, f"P{i}"] for i in range(3)]))
    assert "Mais de 2 linhas" in r.content.decode()
    monkeypatch.setattr(planilhas, "MAX_BYTES", 10)
    r = enviar(admin, LISTA, planilha_csv(["dono", "titulo"], [ana.pk, "Um titulo comprido"]))
    assert "Arquivo maior que" in r.content.decode()


def test_excel_que_explode_ao_abrir_e_recusado(admin, ana, monkeypatch):
    monkeypatch.setattr(planilhas, "_MAX_DESCOMPACTADO", 1000)  # "bomba zip": pequeno fechado, enorme aberto
    r = enviar(admin, LISTA, planilha_xlsx(["dono", "titulo"], [ana.pk, "Cadeira"]), "p.xlsx")
    assert "grande demais depois de aberta" in r.content.decode()


def test_sem_permissao_de_criar_nao_importa(ana):
    ana.is_staff = True
    ana.salvar_como_sistema("teste", update_fields=["is_staff"])
    ana.user_permissions.add(Permission.objects.get(codename="view_pedido"))
    c = Client()
    c.force_login(ana)
    assert c.get(LISTA).status_code == 200
    assert "Importar planilha" not in c.get(LISTA).content.decode()
    assert c.get(LISTA + "importar/").status_code == 403


def test_previa_de_outro_admin_nao_confirma(admin, ana):
    r = enviar(admin, LISTA, planilha_csv(["dono", "titulo"], [ana.pk, "Cadeira"]))
    outro = Usuario.objects.create_superuser("outro@exemplo.com", password=SENHA)
    c = Client()
    c.force_login(outro)
    r2 = c.post(LISTA + "importar/", {"token": token(r), "confirmar": "1"})
    assert r2.status_code == 302 and Pedido.objects.como_sistema("teste").count() == 0


def test_visitante_nao_importa(client):
    r = client.get(LISTA + "importar/")
    assert r.status_code == 302 and "/entrar/" in r["Location"]


# 5. Coluna Acesso
def test_situacao_do_acesso(admin, rf, mailoutbox):
    novo = Usuario.objects.create_user("novo@exemplo.com")
    assert situacao_do_acesso(novo) == "Link não enviado"
    enviar_link_de_senha(rf.get("/"), novo)
    novo.refresh_from_db()
    assert situacao_do_acesso(novo).startswith("Aguardando: link enviado em")
    novo.link_enviado_em = timezone.now() - timedelta(hours=73)
    novo.salvar_como_sistema("teste", update_fields=["link_enviado_em"])
    assert situacao_do_acesso(novo) == "Link vencido, reenviar"
    novo.set_password(SENHA)
    novo.salvar_como_sistema("teste")
    assert situacao_do_acesso(novo) == "Senha definida"


def test_filtro_de_acesso_na_lista(admin, rf, mailoutbox):
    vencido = Usuario.objects.create_user("vencido@exemplo.com")
    vencido.link_enviado_em = timezone.now() - timedelta(days=5)
    vencido.salvar_como_sistema("teste", update_fields=["link_enviado_em"])
    aguardando = Usuario.objects.create_user("aguardando@exemplo.com")
    enviar_link_de_senha(rf.get("/"), aguardando)
    Usuario.objects.create_user("nunca@exemplo.com")
    tela = admin.get(USUARIOS + "?acesso=vencido").content.decode()
    assert "vencido@exemplo.com" in tela and "aguardando@exemplo.com" not in tela and "nunca@exemplo.com" not in tela
    tela = admin.get(USUARIOS + "?acesso=nao_enviado").content.decode()
    assert "nunca@exemplo.com" in tela and "vencido@exemplo.com" not in tela
    tela = admin.get(USUARIOS + "?acesso=aguardando").content.decode()
    assert "aguardando@exemplo.com" in tela and "Aguardando: link enviado em" in tela
    tela = admin.get(USUARIOS + "?acesso=definida").content.decode()
    assert "chefe@exemplo.com" in tela and "nunca@exemplo.com" not in tela


def test_coluna_e_filtro_de_acesso_aparecem_mesmo_com_lista_propria(rf, chefe):
    from django.contrib import admin as dj_admin

    from infra_vibecoding.admin import AdminUsuarioSeguro, FiltroAcesso

    class UsuarioDoSistema(AdminUsuarioSeguro):
        list_display = ("email", "nome", "empresa")
        list_filter = ("is_active",)

    pedido_ = rf.get("/")
    pedido_.user = chefe
    tela = UsuarioDoSistema(Usuario, dj_admin.site)
    assert tela.get_list_display(pedido_) == ["email", "nome", "acesso", "empresa"]
    assert tela.get_list_filter(pedido_)[0] is FiltroAcesso

