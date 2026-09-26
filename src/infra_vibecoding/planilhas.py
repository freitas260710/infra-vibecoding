"""
Importar e exportar planilha na tela de banco (US I.1, decisão D53). Equivalente ao upload/download do "App data"
do Bubble, com as travas do 00.

Toda tabela da tela de banco (AdminSeguro) ganha sozinha os botões "Importar planilha", "Exportar CSV" e
"Exportar Excel". O sistema não escreve nada.

Exportar
- Sai o que está na lista, com a busca e os filtros aplicados. CSV no padrão do Excel brasileiro (ponto e vírgula,
  acentos certos) ou Excel (.xlsx).
- Protegido contra planilha maliciosa: texto que começa com = + - @ ganha um apóstrofo na frente, para o Excel
  não executar como fórmula.
- Nunca sai: senha, chave de sessão, chave do 2FA e códigos de recuperação.

Importar
- Prévia antes de gravar: quantas linhas novas, quantas atualizadas, colunas ignoradas e os erros de cada linha
  (linha, campo, motivo). Só grava quando o admin confirma.
- Tudo ou nada: com qualquer erro, nada é gravado. Corrige a planilha e envia de novo.
- Linha sem "id" cria um registro; com o "id" de um registro que existe, atualiza só as colunas da planilha.
- Relações: a coluna pode trazer o id do registro ligado ("empresa") ou o valor de um campo único dele
  ("empresa__cnpj", "dono__email", "empresa__id_origem"). Serve para dados do Bubble ou de qualquer outra origem.
- Cada linha passa pelas mesmas regras de um cadastro manual (obrigatório, formato, repetido) e é gravada como
  sistema, com registro de quem importou, qual arquivo e em qual tabela.
- Nunca entram: senha, chaves, 2FA, e quem acessa a tela de banco ou é superusuário (permissão se dá uma a uma).
  Usuário importado nasce sem senha e nenhum e-mail sai sozinho: para liberar, use a ação "Enviar link de acesso".
- Até 20.000 linhas e 10 MB por arquivo. O tipo é conferido pelo conteúdo, não pela extensão.
"""
import csv
import datetime
import decimal
import io
import logging
import secrets
import unicodedata
import zipfile

from django.core.cache import cache
from django.core.exceptions import FieldDoesNotExist
from django.db import models, transaction
from django.forms.models import modelform_factory
from django.http import HttpResponse
from django.utils import timezone

log = logging.getLogger("infra_vibecoding.auditoria")

MAX_LINHAS = 20_000
MAX_BYTES = 10 * 1024 * 1024
_MAX_DESCOMPACTADO = 200 * 1024 * 1024  # proteção contra "bomba zip" no .xlsx
_MAX_ERROS_NA_TELA = 200
_VALIDADE_DA_PREVIA = 60 * 60

# Nunca entram nem saem da planilha.
SEGREDOS = frozenset({"password", "chave_de_sessao", "segredo_do_app", "codigos_de_recuperacao",
                      "ultimo_codigo_do_app"})
# Saem na exportação, mas nunca entram pela importação (permissão, datas do login e provas de aceite).
SO_SAIDA = frozenset({"is_superuser", "is_staff", "groups", "user_permissions", "last_login",
                      "email_confirmado_em", "termos_aceitos_em", "dois_fatores", "dois_fatores_desde",
                      "link_enviado_em"})
_PERIGOSOS = ("=", "+", "-", "@", "\t", "\r")
_SIM = {"sim", "s", "x", "true", "verdadeiro", "1", "yes", "y", "on"}
_NAO = {"nao", "n", "false", "falso", "0", "no", "off", ""}


class ErroPlanilha(Exception):
    """Arquivo recusado inteiro (tipo, tamanho, linhas, cabeçalho)."""


def _normal(texto):
    texto = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return "_".join(texto.strip().lower().replace("-", " ").split())


# Leitura do arquivo

def _texto_da_celula(valor):
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "sim" if valor else "não"
    if isinstance(valor, datetime.datetime):
        return valor.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(valor, datetime.date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _ler_xlsx(dados):
    try:
        with zipfile.ZipFile(io.BytesIO(dados)) as z:
            nomes = z.namelist()
            if "xl/workbook.xml" not in nomes:
                raise ErroPlanilha("O arquivo não é uma planilha do Excel (.xlsx).")
            if sum(i.file_size for i in z.infolist()) > _MAX_DESCOMPACTADO:
                raise ErroPlanilha("Planilha grande demais depois de aberta.")
    except zipfile.BadZipFile:
        raise ErroPlanilha("O arquivo não é uma planilha válida.") from None
    import openpyxl

    try:
        livro = openpyxl.load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    except Exception:
        raise ErroPlanilha("Não foi possível abrir a planilha do Excel.") from None
    folha = livro.worksheets[0]
    linhas = []
    for valores in folha.iter_rows(values_only=True):
        linhas.append([_texto_da_celula(v) for v in valores])
        if len(linhas) > MAX_LINHAS + 1:
            break
    livro.close()
    return linhas


def _ler_csv(dados):
    if b"\x00" in dados[:8192]:
        raise ErroPlanilha("O arquivo não é uma planilha (CSV ou Excel).")
    try:
        texto = dados.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = dados.decode("cp1252", errors="strict") if _parece_texto(dados) else None
    if texto is None:
        raise ErroPlanilha("O arquivo não é uma planilha (CSV ou Excel).")
    primeira = texto.split("\n", 1)[0]
    separador = max((";", ",", "\t"), key=primeira.count)
    return [[c.strip() for c in linha] for linha in csv.reader(io.StringIO(texto), delimiter=separador)]


def _parece_texto(dados):
    try:
        dados.decode("cp1252")
    except UnicodeDecodeError:
        return False
    controle = sum(1 for b in dados[:8192] if b < 9 or 13 < b < 32)
    return controle == 0


def ler_arquivo(dados):
    """Devolve (cabeçalho, linhas). Recusa o arquivo inteiro com ErroPlanilha."""
    if len(dados) > MAX_BYTES:
        raise ErroPlanilha(f"Arquivo maior que {MAX_BYTES // (1024 * 1024)} MB. Divida em arquivos menores.")
    if not dados.strip():
        raise ErroPlanilha("O arquivo está vazio.")
    linhas = _ler_xlsx(dados) if dados[:4] == b"PK\x03\x04" else _ler_csv(dados)
    linhas = [linha for linha in linhas if any(c for c in linha)]
    if not linhas:
        raise ErroPlanilha("O arquivo está vazio.")
    cabecalho, corpo = linhas[0], linhas[1:]
    if len(corpo) > MAX_LINHAS:
        raise ErroPlanilha(f"Mais de {MAX_LINHAS} linhas. Divida em arquivos menores.")
    if not corpo:
        raise ErroPlanilha("A planilha só tem o cabeçalho, nenhuma linha de dados.")
    return cabecalho, corpo


# Colunas

def campos_da_exportacao(modelo):
    return [c for c in modelo._meta.concrete_fields if c.name not in SEGREDOS]


def _campos_importaveis(modelo):
    from .arquivos import CampoArquivo

    return {c.name: c for c in modelo._meta.concrete_fields
            if c.editable and not c.primary_key and c.name not in SEGREDOS | SO_SAIDA
            and not isinstance(c, CampoArquivo)  # arquivo não entra por planilha
            and not getattr(c, "auto_now", False) and not getattr(c, "auto_now_add", False)}


class Coluna:
    def __init__(self, posicao, titulo, campo=None, busca=None, eh_id=False):
        self.posicao, self.titulo, self.campo, self.busca, self.eh_id = posicao, titulo, campo, busca, eh_id


def mapear_colunas(modelo, cabecalho):
    """Liga cada coluna da planilha a um campo. Devolve (colunas, ignoradas[(título, motivo)])."""
    importaveis = _campos_importaveis(modelo)
    por_nome = {}
    for nome, campo in importaveis.items():
        for chave in {nome, getattr(campo, "attname", nome), str(campo.verbose_name)}:
            por_nome.setdefault(_normal(chave), campo)
    todos = {_normal(c.name) for c in modelo._meta.concrete_fields} | {_normal(c.verbose_name)
                                                                       for c in modelo._meta.concrete_fields}
    colunas, ignoradas, usados = [], [], set()
    for i, titulo in enumerate(cabecalho):
        chave = _normal(titulo)
        if not chave:
            continue
        if chave == "id" or chave == _normal(modelo._meta.pk.name):
            colunas.append(Coluna(i, titulo, eh_id=True))
            continue
        campo, busca = por_nome.get(chave), None
        if campo is None and "__" in chave:
            base, busca = chave.split("__", 1)
            campo = por_nome.get(base)
            if campo is not None and not campo.is_relation:
                campo = None
        if campo is None:
            if chave in todos or chave.split("__")[0] in {_normal(s) for s in SEGREDOS | SO_SAIDA}:
                ignoradas.append((titulo, "não é importado por segurança"))
            else:
                ignoradas.append((titulo, "não existe nesta tabela"))
            continue
        if campo.name in usados:
            raise ErroPlanilha(f"O campo '{campo.name}' aparece em mais de uma coluna.")
        if busca is not None:
            relacionado = campo.related_model
            try:
                alvo = relacionado._meta.get_field(busca)
            except FieldDoesNotExist:
                raise ErroPlanilha(f"Coluna '{titulo}': '{busca}' não existe em {relacionado._meta.verbose_name}.")
            if not (alvo.unique or alvo.primary_key):
                raise ErroPlanilha(f"Coluna '{titulo}': '{busca}' não é um campo único em "
                                   f"{relacionado._meta.verbose_name}; use o id ou um campo único.")
        usados.add(campo.name)
        colunas.append(Coluna(i, titulo, campo=campo, busca=busca))
    if not any(c.campo for c in colunas):
        raise ErroPlanilha("Nenhuma coluna da planilha corresponde a um campo desta tabela.")
    return colunas, ignoradas


# Conversão de cada célula para o formulário do Django

def _valor_para_formulario(campo, texto):
    if texto.startswith("'") and texto[1:2] in _PERIGOSOS:
        texto = texto[1:]  # desfaz a proteção contra fórmula da exportação
    if isinstance(campo, models.BooleanField):
        chave = _normal(texto)
        if chave in _SIM:
            return "on"
        if chave in _NAO:
            return ""
        raise ValueError("Use sim ou não.")
    if isinstance(campo, (models.DecimalField, models.FloatField)) and "," in texto:
        return texto.replace(".", "").replace(",", ".")
    if campo.choices and texto:
        for valor, rotulo in campo.flatchoices:
            if _normal(texto) == _normal(rotulo) and str(valor) != texto:
                return str(valor)
    return texto


class Resultado:
    def __init__(self):
        self.novos = 0
        self.atualizados = 0
        self.erros = []       # (linha, coluna, mensagem)
        self.ignoradas = []   # (título, motivo)
        self.formularios = []
        self.amostra = []
        self.total_de_erros = 0

    @property
    def ok(self):
        return self.total_de_erros == 0

    def erro(self, linha, coluna, mensagem):
        self.total_de_erros += 1
        if len(self.erros) < _MAX_ERROS_NA_TELA:
            self.erros.append((linha, coluna, mensagem))


def conferir(modelo, cabecalho, linhas, pode_criar, pode_atualizar):
    """Confere todas as linhas sem gravar nada. Deve rodar dentro de uma tela da tela de banco."""
    resultado = Resultado()
    colunas, resultado.ignoradas = mapear_colunas(modelo, cabecalho)
    campos = [c for c in colunas if c.campo]
    coluna_id = next((c for c in colunas if c.eh_id), None)
    Formulario = modelform_factory(modelo, fields=[c.campo.name for c in campos])
    unicos = [c.campo.name for c in campos if c.campo.unique]
    vistos = {nome: {} for nome in unicos}
    cache_busca = {}
    presentes = {c.campo.name for c in campos}
    faltando = [c.name for c in _campos_importaveis(modelo).values()
                if c.name not in presentes and not c.blank and not c.has_default() and not c.null]

    for n, linha in enumerate(linhas, start=2):  # linha 1 é o cabeçalho
        valores = linha + [""] * (len(cabecalho) - len(linha))
        instancia = None
        if coluna_id is not None and valores[coluna_id.posicao].strip():
            pk = valores[coluna_id.posicao].strip()
            try:
                instancia = modelo._base_manager.filter(pk=pk).first()
            except (ValueError, TypeError):
                instancia = None
            if instancia is None:
                resultado.erro(n, coluna_id.titulo, f"Não existe registro com id {pk} nesta tabela.")
                continue
            if not pode_atualizar:
                resultado.erro(n, coluna_id.titulo, "Você não tem permissão para alterar registros desta tabela.")
                continue
        elif not pode_criar:
            resultado.erro(n, "", "Você não tem permissão para criar registros nesta tabela.")
            continue

        if instancia is None and faltando:
            resultado.erro(n, "", "Para criar, a planilha precisa da coluna: " + ", ".join(faltando) + ".")
            continue
        dados, com_erro = {}, set()
        for coluna in campos:
            texto = valores[coluna.posicao] if coluna.posicao < len(valores) else ""
            try:
                valor = _valor_para_formulario(coluna.campo, texto)
            except ValueError as e:
                resultado.erro(n, coluna.titulo, str(e))
                com_erro.add(coluna.campo.name)
                continue
            if coluna.busca and valor:
                chave = (coluna.campo.name, valor)
                if chave not in cache_busca:
                    relacionado = coluna.campo.related_model
                    alvo = relacionado._base_manager.filter(**{coluna.busca: valor}).values_list("pk", flat=True)
                    cache_busca[chave] = next(iter(alvo[:1]), None)
                if cache_busca[chave] is None:
                    resultado.erro(n, coluna.titulo, f"Não encontrado: {valor}.")
                    com_erro.add(coluna.campo.name)
                    continue
                valor = str(cache_busca[chave])
            dados[coluna.campo.name] = valor
        # Campos sem coluna na planilha: numa atualização continuam como estão; num registro novo, o padrão.
        formulario = Formulario(data=dados, instance=instancia)
        valido = formulario.is_valid()
        for campo, mensagens in ({} if valido else formulario.errors).items():
            if campo in com_erro:
                continue  # já apontado acima (ex.: valor de sim/não inválido)
            titulo = next((c.titulo for c in campos if c.campo.name == campo), campo)
            for mensagem in mensagens:
                resultado.erro(n, "" if campo == "__all__" else titulo, mensagem)
        if not valido or com_erro:
            continue
        repetida = False
        for nome in unicos:
            valor = formulario.cleaned_data.get(nome)
            if valor in (None, ""):
                continue
            chave = valor.lower() if isinstance(valor, str) else valor
            if chave in vistos[nome]:
                titulo = next(c.titulo for c in campos if c.campo.name == nome)
                resultado.erro(n, titulo, f"Valor repetido na planilha (já está na linha {vistos[nome][chave]}).")
                repetida = True
            else:
                vistos[nome][chave] = n
        if repetida:
            continue
        if instancia is None:
            resultado.novos += 1
        else:
            resultado.atualizados += 1
        resultado.formularios.append((n, formulario))
        if len(resultado.amostra) < 10:
            resultado.amostra.append([n, "atualizar" if instancia else "criar",
                                      *[_texto_da_celula(formulario.cleaned_data.get(c.campo.name))
                                        for c in campos]])
    resultado.titulos_da_amostra = [c.titulo for c in campos]
    return resultado


def gravar(resultado, motivo):
    """Grava tudo numa transação só: se qualquer linha falhar, nada fica gravado."""
    with transaction.atomic():
        for n, formulario in resultado.formularios:
            obj = formulario.save(commit=False)
            obj.salvar_como_sistema(f"{motivo} (linha {n})")
            formulario.save_m2m()


# Guardar o arquivo entre a prévia e a confirmação (1 hora)

def guardar_para_confirmar(dados, nome, modelo, usuario):
    token = secrets.token_urlsafe(24)
    cache.set(f"infra_vibecoding:planilha:{token}", {
        "dados": dados, "nome": nome, "modelo": modelo._meta.label, "usuario": usuario.pk,
    }, _VALIDADE_DA_PREVIA)
    return token


def recuperar(token, modelo, usuario):
    guardado = cache.get(f"infra_vibecoding:planilha:{token}") if token else None
    if not guardado or guardado["modelo"] != modelo._meta.label or guardado["usuario"] != usuario.pk:
        return None
    return guardado


def descartar(token):
    cache.delete(f"infra_vibecoding:planilha:{token}")


# Exportar

def _protegido(texto):
    return "'" + texto if texto.startswith(_PERIGOSOS) else texto


def _valor_exportado(campo, obj, para_excel):
    valor = getattr(obj, campo.attname)
    if valor is None:
        return ""
    if isinstance(valor, bool):
        return "sim" if valor else "não"
    if isinstance(valor, datetime.datetime):
        if timezone.is_aware(valor):
            valor = timezone.localtime(valor).replace(tzinfo=None)
        return valor if para_excel else valor.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(valor, datetime.date):
        return valor if para_excel else valor.strftime("%d/%m/%Y")
    if isinstance(valor, (int, float, decimal.Decimal)) and not isinstance(valor, bool):
        if para_excel:
            return valor
        return str(valor).replace(".", ",") if isinstance(valor, (float, decimal.Decimal)) else str(valor)
    return _protegido(str(valor))


def exportar(modelo, queryset, formato, nome_base):
    campos = campos_da_exportacao(modelo)
    cabecalho = [c.name for c in campos]
    data = timezone.localdate().strftime("%Y-%m-%d")
    if formato == "xlsx":
        import openpyxl

        livro = openpyxl.Workbook(write_only=True)
        folha = livro.create_sheet(title=str(modelo._meta.verbose_name_plural)[:31] or "dados")
        folha.append(cabecalho)
        total = 0
        for obj in queryset.iterator():
            folha.append([_valor_exportado(c, obj, True) for c in campos])
            total += 1
        saida = io.BytesIO()
        livro.save(saida)
        resposta = HttpResponse(
            saida.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resposta["Content-Disposition"] = f'attachment; filename="{nome_base}-{data}.xlsx"'
        return resposta, total
    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";", lineterminator="\r\n")
    escritor.writerow(cabecalho)
    total = 0
    for obj in queryset.iterator():
        escritor.writerow([_valor_exportado(c, obj, False) for c in campos])
        total += 1
    resposta = HttpResponse(("﻿" + saida.getvalue()).encode("utf-8"), content_type="text/csv; charset=utf-8")
    resposta["Content-Disposition"] = f'attachment; filename="{nome_base}-{data}.csv"'
    return resposta, total
